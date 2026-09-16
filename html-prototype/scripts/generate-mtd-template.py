#!/usr/bin/env python3
"""
Generates the FreeAgent MTD quarterly spreadsheet template.

Structure and category vocabulary now follow the team's own sample
("Sample MTD Template for IT.xlsx"): three tabs only — Sales transactions,
Purchase transactions, Quarterly filing codes. No separate Information or
Quarterly filing totals tab; category totals are computed by us once we
parse the transaction rows, not by in-sheet formulas.

The category list below is FreeAgent's own analysis vocabulary (not the
HMRC SA103 box wording used in an earlier draft of this template), taken
from the 'Quarterly filing codes' tab of the sample, with obvious typos/
stray whitespace cleaned up (e.g. "Disallowable Expneses " -> "Disallowable
Expenses"). Keep this list in sync with whatever FreeAgent nominal-code
mapping the processing pipeline ends up using.

Run: python3 generate-mtd-template.py
Output: ../assets/mtd-spreadsheet-template.xlsx
"""

import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

# Source of truth for the "Quarterly filing analysis" dropdowns.
# Order matters: the Sales dropdown reads the income block, the Purchase
# dropdown reads the expense block (see codes_first_income_row / codes_first_expense_row below).
INCOME_CATEGORIES = [
    "Sales",
    "CIS Income",
]
INCOME_EXCLUDE_LABEL = "Disallowable Income"

EXPENSE_CATEGORIES = [
    "Cost of Sales",
    "CIS Deductions",
    "Net Salary Expense",
    "Motor Expenses",
    "Rent",
    "Office Costs",
    "Internet & Telephone",
    "Advertising and Promotion",
    "Business Entertaining",
    "Interest Payable",
    "Bank/Finance Charges",
    "Bad Debts Written Off",
    "Accountancy Fees",
    "Depreciation",
    "Depreciation Charge",
]
EXPENSE_EXCLUDE_LABEL = "Disallowable Expenses"

PRIMARY_DARK = "12508F"
TEXT_MUTED = "5B6B7C"
GREY = "F2F4F7"
WARNING = "B3261E"

TX_DATA_ROWS = 191  # rows 8 to 198
TX_FIRST_ROW = 8
TX_LAST_ROW = TX_FIRST_ROW + TX_DATA_ROWS - 1  # 198
INSERT_WARNING_ROW = TX_LAST_ROW + 1  # 199

thin = Side(style="thin", color="D5DCE3")
border_all = Border(left=thin, right=thin, top=thin, bottom=thin)
muted_italic = Font(italic=True, color=TEXT_MUTED, size=10)


def title_bar(ws, row, text, span_cols, size=13):
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(size=size, bold=True, color="FFFFFF")
    cell.alignment = Alignment(vertical="center")
    for col in range(1, span_cols + 1):
        ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor=PRIMARY_DARK)
    ws.row_dimensions[row].height = 24
    if span_cols > 1:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span_cols)


def build_codes_sheet(wb):
    ws = wb.create_sheet("Quarterly filing codes")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 32

    title_bar(ws, 2, "MTD Quarterly filing analysis for quarterly income and expenditure updates", 2)
    ws.merge_cells("A2:B2")

    note = ws.cell(
        row=3, column=2,
        value=(
            "This information is used to populate the Quarterly filing analysis dropdown on the "
            "Sales transactions and Purchase transactions tabs. Businesses with a turnover of less "
            "than £90,000 can still list a single line per category if they don't want to log every "
            "individual transaction — the figures still come from these tabs."
        ),
    )
    note.font = muted_italic
    note.alignment = Alignment(wrap_text=True)
    ws.row_dimensions[3].height = 56

    ws.cell(row=5, column=2, value="Income/Expense analysis").font = Font(bold=True, color=PRIMARY_DARK)

    row = 6
    for label in INCOME_CATEGORIES:
        ws.cell(row=row, column=2, value=label)
        row += 1
    ws.cell(row=row, column=2, value=INCOME_EXCLUDE_LABEL)
    row += 1
    codes_first_expense_row = row
    for label in EXPENSE_CATEGORIES:
        ws.cell(row=row, column=2, value=label)
        row += 1
    ws.cell(row=row, column=2, value=EXPENSE_EXCLUDE_LABEL)

    return ws


def build_transactions_sheet(wb, sheet_name, party_label, ref_label, date_paid_label,
                              dropdown_range, example_rows, table_name):
    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False
    widths = {"A": 13, "B": 16, "C": 22, "D": 26, "E": 26, "F": 13, "G": 24, "H": 13}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    title_bar(ws, 2, sheet_name, 8)

    kind = "sales" if "Sales" in sheet_name else "purchase"
    ws.cell(row=3, column=1, value=f"Enter or import all {kind} transactions")
    ws.cell(
        row=4, column=1,
        value=(
            "Please enter a Quarterly filing analysis for all transactions (unless you wish to "
            "exclude an item from this quarterly Income Tax filing)"
        ),
    )
    ws.cell(
        row=5, column=1,
        value=(
            f"The totals are calculated from amounts in rows {TX_FIRST_ROW} to {TX_LAST_ROW}. Please "
            "insert rows within these row numbers to ensure amounts continue to be included in the totals."
        ),
    )
    for r in (3, 4, 5):
        ws.cell(row=r, column=1).font = muted_italic
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)

    headers = ["Invoice date", ref_label, party_label, "Description", "Quarterly filing analysis",
               "Amount (£)", "Comments", date_paid_label]
    header_row = 6
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=PRIMARY_DARK)
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[header_row].height = 30

    do_not_enter_row = header_row + 1  # 7
    ws.cell(row=do_not_enter_row, column=2, value="[DO NOT ENTER ANY DATA IN THIS ROW >>]").font = Font(italic=True, color=TEXT_MUTED)
    total_cell = ws.cell(row=do_not_enter_row, column=6, value=f"=SUM(F{TX_FIRST_ROW}:F{TX_LAST_ROW})")
    total_cell.font = Font(bold=True)
    total_cell.number_format = "#,##0.00;[Red]-#,##0.00"
    for col in range(1, 9):
        ws.cell(row=do_not_enter_row, column=col).fill = PatternFill("solid", fgColor=GREY)

    for i, row_data in enumerate(example_rows):
        r = TX_FIRST_ROW + i
        for col, val in enumerate(row_data, start=1):
            ws.cell(row=r, column=col, value=val)

    for r in range(TX_FIRST_ROW, TX_LAST_ROW + 1):
        ws.cell(row=r, column=1).number_format = "dd/mm/yyyy"
        ws.cell(row=r, column=6).number_format = "#,##0.00;[Red]-#,##0.00"
        ws.cell(row=r, column=8).number_format = "dd/mm/yyyy"
        for col in range(1, 9):
            ws.cell(row=r, column=col).border = border_all

    warning = ws.cell(row=INSERT_WARNING_ROW, column=1, value="IF YOU NEED TO INSERT MORE ROWS, ONLY DO SO ABOVE THIS ROW")
    warning.font = Font(bold=True, color=WARNING, size=9)
    ws.merge_cells(start_row=INSERT_WARNING_ROW, start_column=1, end_row=INSERT_WARNING_ROW, end_column=8)

    table_ref = f"A{header_row}:H{TX_LAST_ROW}"
    table = Table(displayName=table_name, ref=table_ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(table)

    dv = DataValidation(type="list", formula1=dropdown_range, allow_blank=True, showDropDown=False)
    dv.error = "Please choose a category from the list"
    dv.errorTitle = "Invalid category"
    ws.add_data_validation(dv)
    for r in range(TX_FIRST_ROW, TX_LAST_ROW + 1):
        dv.add(ws.cell(row=r, column=5))

    ws.freeze_panes = f"A{header_row + 1}"
    return ws


def main():
    wb = Workbook()
    wb.remove(wb.active)

    codes_first_income_row = 6
    codes_first_expense_row = codes_first_income_row + len(INCOME_CATEGORIES) + 1  # +1 for the income exclude row
    sales_dropdown_last_row = codes_first_expense_row - 1
    purchase_dropdown_last_row = codes_first_expense_row + len(EXPENSE_CATEGORIES)  # includes expense exclude row

    build_transactions_sheet(
        wb, "Sales transactions",
        party_label="Customer", ref_label="Invoice No/Ref.", date_paid_label="Date paid",
        dropdown_range=f"='Quarterly filing codes'!$B${codes_first_income_row}:$B${sales_dropdown_last_row}",
        example_rows=[
            ("06/04/2026", "INV0001", "Customer A", "Consulting invoice", "Sales", 500.00, "", ""),
            ("14/04/2026", "INV0002", "Customer B", "Retainer", "CIS Income", 400.00, "", ""),
            ("20/04/2026", "N/A", "N/A", "Capital introduced", "Disallowable Income", 6000.00, "Excluded from Income Tax filing", ""),
        ],
        table_name="SalesTransactions",
    )
    build_transactions_sheet(
        wb, "Purchase transactions",
        party_label="Supplier", ref_label="Reference (optional)", date_paid_label="Date paid\n(if different)",
        dropdown_range=f"='Quarterly filing codes'!$B${codes_first_expense_row}:$B${purchase_dropdown_last_row}",
        example_rows=[
            ("09/04/2026", "", "Supplier A", "Stock purchase", "Cost of Sales", -100.00, "", ""),
            ("14/04/2026", "", "Supplier B", "Software subscription", "Office Costs", -45.00, "", ""),
            ("20/04/2026", "", "N/A", "Drawings", "Disallowable Expenses", -1600.00, "Excluded from Income Tax filing", ""),
        ],
        table_name="PurchaseTransactions",
    )
    build_codes_sheet(wb)

    wb.active = 0

    out_dir = os.path.join(os.path.dirname(__file__), "..", "assets")
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, "mtd-spreadsheet-template.xlsx")
    wb.save(out_path)
    print(f"Wrote {os.path.abspath(out_path)}")


if __name__ == "__main__":
    main()
